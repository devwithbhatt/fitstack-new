from django.conf import settings
import logging
from django.shortcuts import render, redirect, get_object_or_404
from django.http import JsonResponse
from django.views.decorators.http import require_POST
from .models import Enquiry
from .forms import EnquiryForm
from django.contrib import messages
from django.core.paginator import Paginator
from django.contrib.auth.decorators import login_required
from apps.login.decorators import custom_permission_required
from django.views.decorators.cache import never_cache
from django.db.models import Q, F, Case, When, Value, IntegerField
from django.utils import timezone
from apps.whatsapp.services import WhatsAppService

# Get an instance of a logger
logger = logging.getLogger(__name__)

@never_cache
@login_required
@custom_permission_required('change_enquiry')
def update_enquiry_status(request, enquiry_id):
    gym = getattr(request, 'gym', None)
    enquiry = Enquiry.objects.filter(id=enquiry_id, gym=gym).first()
    if not enquiry:
        messages.error(request, 'Enquiry not found.')
        return redirect('enquiry_list')
    if request.method == 'POST':
        status = request.POST.get('status')
        if status in [choice[0] for choice in Enquiry.STATUS_CHOICES]:
            enquiry.status = status
            enquiry.save()
            messages.success(request, f"Status for {enquiry.name} updated successfully.")
    return redirect('enquiry_list')

@never_cache
@login_required
@custom_permission_required('add_enquiry')
def add_new_enquiry(request):
    gym = getattr(request, 'gym', None)
    if request.method == 'POST':
        form = EnquiryForm(request.POST)
        if form.is_valid():
            enquiry = form.save(commit=False)
            enquiry.gym = gym
            enquiry.save()

            # Send WhatsApp confirmation via the centralized service
            if gym and gym.whatsapp_enabled:
                try:
                    logo_url = gym.logo.url if gym and gym.logo else None

                    whatsapp_service = WhatsAppService(gym_id=gym.id)
                    result = whatsapp_service.send_enquiry_confirmation(
                        request=request,
                        to_number=enquiry.mobile_number,
                        name=enquiry.name,
                        gym_name=gym.name if gym else "FitStack Gym Management",
                        gym_contact_number=gym.phone if gym else "N/A",
                        gym_logo_url=logo_url
                    )
                    if result.get('success'):
                        messages.success(request, 'Enquiry added and confirmation sent!')
                    else:
                        # Log the detailed error for debugging
                        error_details = result.get('error', 'Unknown error')
                        logger.error(f"WhatsApp API failed for enquiry {enquiry.id}: {error_details}")
                        messages.warning(request, "Enquiry added, but the confirmation message could not be sent.")
                except Exception as e:
                    logger.error(f"An unexpected error occurred sending WhatsApp confirmation for enquiry {enquiry.id}: {e}")
                    messages.error(request, "Enquiry added, but an unexpected error occurred while sending the confirmation.")
            else:
                messages.success(request, 'Enquiry added successfully!')
            
            return redirect('enquiry_list')
    else:
        form = EnquiryForm()
    return render(request, 'enquiry/add_new_enquiry.html', {'form': form})

@never_cache
@login_required
@custom_permission_required('view_enquiry')
def enquiry_list(request):
    gym = getattr(request, 'gym', None)
    today = timezone.localdate()
    enquiry_list = Enquiry.objects.filter(gym=gym).annotate(
        date_sort=Case(
            When(next_follow_up_date__gte=today, then=Value(0)),
            When(next_follow_up_date__isnull=True, then=Value(1)),
            When(next_follow_up_date__lt=today, then=Value(2)),
            default=Value(3),
            output_field=IntegerField(),
        )
    ).order_by('date_sort', 'next_follow_up_date')
    
    # Search functionality
    query = request.GET.get('q')
    status_filter = request.GET.get('status')
    date_from = request.GET.get('date_from')
    date_to = request.GET.get('date_to')

    if query:
        enquiry_list = enquiry_list.filter(
            Q(name__icontains=query) |
            Q(mobile_number__icontains=query) |
            Q(email__icontains=query)
        ).distinct()

    if status_filter:
        enquiry_list = enquiry_list.filter(status=status_filter)

    if date_from:
        enquiry_list = enquiry_list.filter(next_follow_up_date__gte=date_from)

    if date_to:
        enquiry_list = enquiry_list.filter(next_follow_up_date__lte=date_to)

    if request.GET.get('export') == 'excel':
        import openpyxl
        from openpyxl.styles import Font, Alignment
        from django.http import HttpResponse

        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename="enquiries.xlsx"'

        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.title = 'Enquiries'

        # Write header
        columns = [
            'SR. No.', 'Name', 'Contact Number', 'Interested In', 
            'Enquiry Date', 'Follow-up Date', 'Status', 'Remark'
        ]
        
        for col_num, column_title in enumerate(columns, 1):
            cell = worksheet.cell(row=1, column=col_num)
            cell.value = column_title
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            column_letter = openpyxl.utils.get_column_letter(col_num)
            worksheet.column_dimensions[column_letter].width = 20

        # Write data rows
        for index, enquiry in enumerate(enquiry_list, 1):
            row = [
                index,
                enquiry.name,
                enquiry.mobile_number,
                enquiry.get_interested_in_display() if hasattr(enquiry, 'get_interested_in_display') else '',
                enquiry.enquiry_date.strftime('%d-%m-%Y') if enquiry.enquiry_date else '',
                enquiry.next_follow_up_date.strftime('%d-%m-%Y') if enquiry.next_follow_up_date else '',
                enquiry.get_status_display() if hasattr(enquiry, 'get_status_display') else '',
                enquiry.enquiry_note or ''
            ]
            for col_num, cell_value in enumerate(row, 1):
                cell = worksheet.cell(row=index + 1, column=col_num)
                cell.value = cell_value
                cell.alignment = Alignment(vertical='center')

        workbook.save(response)
        return response

    # Pagination
    paginator = Paginator(enquiry_list, 10)  # Show 10 enquiries per page
    page_number = request.GET.get('page')
    enquiries = paginator.get_page(page_number)

    return render(request, 'enquiry/enquiry_list.html', {
        'enquiries': enquiries,
        'status_choices': Enquiry.STATUS_CHOICES,
    })

@never_cache
@login_required
@custom_permission_required('change_enquiry')
def edit_enquiry(request, enquiry_id):
    gym = getattr(request, 'gym', None)
    enquiry = Enquiry.objects.filter(id=enquiry_id, gym=gym).first()
    if not enquiry:
        messages.error(request, 'Enquiry not found.')
        return redirect('enquiry_list')
    if request.method == 'POST':
        form = EnquiryForm(request.POST, instance=enquiry)
        if form.is_valid():
            form.save()
            messages.success(request, 'Enquiry updated successfully!')
            return redirect('enquiry_list')
    else:
        form = EnquiryForm(instance=enquiry)
    return render(request, 'enquiry/edit_enquiry.html', {'form': form})

@never_cache
@login_required
@require_POST
@custom_permission_required('delete_enquiry')
def delete_enquiry(request, enquiry_id):
    gym = getattr(request, 'gym', None)
    enquiry = Enquiry.objects.filter(id=enquiry_id, gym=gym).first()
    if not enquiry:
        return JsonResponse({'status': 'error', 'message': 'Enquiry not found.'}, status=404)
    try:
        enquiry.delete()
        messages.success(request, 'Enquiry has been deleted successfully.')
        return JsonResponse({'status': 'success', 'message': 'Enquiry deleted successfully.'})
    except Exception as e:
        messages.error(request, f'An error occurred: {e}')
        return JsonResponse({'status': 'error', 'message': str(e)}, status=500)